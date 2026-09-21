from io import BytesIO

from docx import Document

from geoagent.report.template import parse_docx
from geoagent.report.template_fill import refill_docx


def test_refill_keeps_paragraph_run_style_and_fills_table(tmp_path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('总数').bold = True
    paragraph.add_run('{{total_n}}')
    paragraph.add_run('个')
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = '地区'
    table.cell(0, 1).text = '面积'
    source = BytesIO()
    document.save(source)
    template = tmp_path / 'template.docx'
    template.write_bytes(source.getvalue())
    parsed = parse_docx(source.getvalue(), template.name)
    scalar = next(slot for slot in parsed['slots'] if slot['kind'] == 'placeholder')
    table_slot = next(slot for slot in parsed['slots'] if slot['kind'] == 'table_region')
    output = tmp_path / 'filled.docx'
    refill_docx(
        template,
        parsed,
        output,
        scalar_answers={scalar['id']: 12},
        table_answers={table_slot['id']: [['杭州市', '1.20'], ['宁波市', '0.80']]},
        text_replacements={'总数': '本期总数'},
    )
    result = Document(output)
    assert result.paragraphs[0].text == '本期总数12个'
    assert result.paragraphs[0].runs[0].bold is True
    assert result.tables[0].cell(1, 0).text == '杭州市'
